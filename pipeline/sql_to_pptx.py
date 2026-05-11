"""
SQL Server -> Excel -> PowerPoint weekly report pipeline.
Run directly or via run_pipeline.bat (Windows Task Scheduler).
Config is loaded from config.env in the same directory.
"""

import os
import sys
import logging
from datetime import datetime
from pathlib import Path

import pyodbc
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# Bootstrap
# ---------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).parent.resolve()
load_dotenv(SCRIPT_DIR / "config.env")

LOG_FILE = SCRIPT_DIR / "pipeline.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Config (loaded from config.env; all values have safe defaults for testing)
# ---------------------------------------------------------------------------

SQL_SERVER   = os.getenv("SQL_SERVER",   "localhost")
SQL_DATABASE = os.getenv("SQL_DATABASE", "YourDatabase")
SQL_USERNAME = os.getenv("SQL_USERNAME", "")          # leave blank for Windows Auth
SQL_PASSWORD = os.getenv("SQL_PASSWORD", "")
SQL_DRIVER   = os.getenv("SQL_DRIVER",   "ODBC Driver 17 for SQL Server")
SQL_QUERY    = os.getenv("SQL_QUERY",    "SELECT TOP 50 * FROM dbo.YourTable")

OUTPUT_DIR   = Path(os.getenv("OUTPUT_DIR", str(SCRIPT_DIR / "output")))

# Slide display cap — rows beyond this still appear in Excel
PPT_MAX_ROWS = int(os.getenv("PPT_MAX_ROWS", "25"))

# Branding colours (hex without #)
COLOR_HEADER_BG   = RGBColor(0x1F, 0x4E, 0x79)   # dark navy
COLOR_HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_ROW_ALT     = RGBColor(0xD6, 0xE4, 0xF0)   # light blue stripe
COLOR_ROW_BASE    = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_SLIDE_BG    = RGBColor(0xF2, 0xF7, 0xFF)
COLOR_TITLE_TEXT  = RGBColor(0x1F, 0x4E, 0x79)
COLOR_SUBTITLE    = RGBColor(0x70, 0x70, 0x70)
COLOR_FOOTER      = RGBColor(0x99, 0x99, 0x99)

# ---------------------------------------------------------------------------
# SQL
# ---------------------------------------------------------------------------

def _build_conn_string() -> str:
    base = f"DRIVER={{{SQL_DRIVER}}};SERVER={SQL_SERVER};DATABASE={SQL_DATABASE};"
    if SQL_USERNAME and SQL_PASSWORD:
        return base + f"UID={SQL_USERNAME};PWD={SQL_PASSWORD};"
    return base + "Trusted_Connection=yes;"


def fetch_data() -> pd.DataFrame:
    log.info("Connecting to SQL Server: %s / %s", SQL_SERVER, SQL_DATABASE)
    conn = pyodbc.connect(_build_conn_string(), timeout=30)
    try:
        df = pd.read_sql(SQL_QUERY, conn)
    finally:
        conn.close()
    log.info("Fetched %d rows x %d columns.", len(df), len(df.columns))
    return df

# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_THIN_SIDE  = Side(style="thin",   color="CCCCCC")
_CELL_BORDER = Border(
    left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE
)


def export_to_excel(df: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Report"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Header row
    for col_i, col_name in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=col_i, value=col_name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
        c.border = _CELL_BORDER

    ws.row_dimensions[1].height = 20

    # Data rows
    for row_i, row in enumerate(df.itertuples(index=False), start=2):
        alt = row_i % 2 == 0
        fill_color = "D6E4F0" if alt else "FFFFFF"
        row_fill = PatternFill("solid", fgColor=fill_color)
        for col_i, value in enumerate(row, start=1):
            display = "" if value is None else value
            c = ws.cell(row=row_i, column=col_i, value=display)
            c.fill = row_fill
            c.font = Font(size=10, name="Calibri")
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = _CELL_BORDER

    # Auto-size columns (cap at 45 chars wide)
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(width + 3, 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    log.info("Excel saved: %s", path)


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------

def _set_cell(cell, text: str, bg: RGBColor, font_color: RGBColor,
              bold: bool = False, font_size: int = 9,
              align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = bg
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.text = text
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "Calibri"


def build_pptx(df: pd.DataFrame, path: Path, excel_filename: str) -> None:
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Slide background
    bg_fill = slide.background.fill
    bg_fill.solid()
    bg_fill.fore_color.rgb = COLOR_SLIDE_BG

    # --- Title ---
    title_box = slide.shapes.add_textbox(Inches(0.35), Inches(0.18), Inches(12.6), Inches(0.65))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Weekly Report  —  {datetime.now().strftime('%B %d, %Y')}"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE_TEXT
    run.font.name = "Calibri"

    # --- Subtitle ---
    sub_box = slide.shapes.add_textbox(Inches(0.35), Inches(0.83), Inches(12.6), Inches(0.32))
    stf = sub_box.text_frame
    sp = stf.paragraphs[0]
    srun = sp.add_run()
    srun.text = (
        f"Source: SQL Server • Database: {SQL_DATABASE} • "
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    srun.font.size = Pt(9)
    srun.font.color.rgb = COLOR_SUBTITLE
    srun.font.name = "Calibri"

    # --- Divider ---
    divider = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.35), Inches(1.17), Inches(12.6), Inches(0.04)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = COLOR_HEADER_BG
    divider.line.fill.background()

    # --- Table ---
    display_df  = df.head(PPT_MAX_ROWS)
    n_data_rows = len(display_df)
    n_cols      = len(display_df.columns)

    tbl = slide.shapes.add_table(
        n_data_rows + 1, n_cols,
        Inches(0.35), Inches(1.28),
        Inches(12.6), Inches(5.95)
    ).table

    # Even column widths
    col_width = int(Inches(12.6) / n_cols)
    for i in range(n_cols):
        tbl.columns[i].width = col_width

    # Header row
    for col_i, col_name in enumerate(display_df.columns):
        _set_cell(
            tbl.cell(0, col_i), str(col_name),
            bg=COLOR_HEADER_BG, font_color=COLOR_HEADER_TEXT,
            bold=True, font_size=10, align=PP_ALIGN.CENTER
        )

    # Data rows
    for row_i, row in enumerate(display_df.itertuples(index=False), start=1):
        bg = COLOR_ROW_ALT if row_i % 2 == 0 else COLOR_ROW_BASE
        for col_i, value in enumerate(row):
            _set_cell(
                tbl.cell(row_i, col_i),
                "" if value is None else str(value),
                bg=bg, font_color=RGBColor(0x20, 0x20, 0x20),
                font_size=9
            )

    # --- Footer ---
    footer_text = (
        f"Showing {n_data_rows} of {len(df)} rows  —  "
        f"Full dataset: {excel_filename}"
        if len(df) > PPT_MAX_ROWS
        else f"Full dataset: {excel_filename}"
    )
    footer_box = slide.shapes.add_textbox(Inches(0.35), Inches(7.22), Inches(12.6), Inches(0.22))
    ftf = footer_box.text_frame
    fp = ftf.paragraphs[0]
    frun = fp.add_run()
    frun.text = footer_text
    frun.font.size = Pt(7.5)
    frun.font.color.rgb = COLOR_FOOTER
    frun.font.name = "Calibri"

    prs.save(path)
    log.info("PowerPoint saved: %s", path)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def run() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    log.info("=" * 60)
    log.info("Pipeline started")

    df = fetch_data()

    excel_path = OUTPUT_DIR / f"report_{stamp}.xlsx"
    export_to_excel(df, excel_path)

    pptx_path = OUTPUT_DIR / f"report_{stamp}.pptx"
    build_pptx(df, pptx_path, excel_path.name)

    log.info("Pipeline complete.  Excel: %s  |  PPTX: %s", excel_path.name, pptx_path.name)
    log.info("=" * 60)


if __name__ == "__main__":
    run()
